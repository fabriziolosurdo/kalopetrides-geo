"""Minimal Markdown -> DOCX converter (python-docx) for the client report: headings, paragraphs,
**bold**, *italic*, numbered/bulleted lists, pipe tables. Sober formatting, no logo."""
import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
src, dst = sys.argv[1], sys.argv[2]
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2.2)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
def runs(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part: continue
        if part.startswith("**"): par.add_run(part[2:-2]).bold = True
        elif part.startswith("*"): par.add_run(part[1:-1]).italic = True
        else: par.add_run(part)
lines = open(src, encoding="utf-8").read().splitlines(); i = 0
while i < len(lines):
    l = lines[i]
    if not l.strip(): i += 1; continue
    if l.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                p = t.cell(r, c).paragraphs[0]; runs(p, val)
                if r == 0:
                    for run in p.runs: run.bold = True
        doc.add_paragraph(); continue
    m = re.match(r"^(#{1,3})\s+(.*)", l)
    if m: doc.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0); i += 1; continue
    m = re.match(r"^\d+\.\s+(.*)", l)
    if m: runs(doc.add_paragraph(style="List Number"), m.group(1)); i += 1; continue
    m = re.match(r"^[-*]\s+(.*)", l)
    if m: runs(doc.add_paragraph(style="List Bullet"), m.group(1)); i += 1; continue
    runs(doc.add_paragraph(), l); i += 1
doc.save(dst); print("saved", dst)
